from __future__ import annotations

import datetime as dt
import io
from collections.abc import Callable

import pytest

from ezmd.ir import Comment, Document, Heading, Image, ListBlock, Paragraph, Slide, Table, WarningKind
from ezmd_converters.office.pptx import PptxConverter
from ezmd_converters.office.xlsx import XlsxConverter
from ezmd_converters.office.xlsx_cells import XCell, a1, col_letter, fmt_value, header_rows, regions

Run = Callable[..., Document]


def test_pptx_slides_notes_tables_chart(run: Run, fixture_bytes: Callable[[str], bytes]) -> None:
    doc = run(PptxConverter(), fixture_bytes("pptx-lecture"), "deck.pptx")
    slides = [b for b in doc.blocks if isinstance(b, Slide)]
    assert [s.title for s in slides] == [
        "Pond Ecology 101",
        "Why ponds matter",
        "Survey results",
        "Species by season",
        "Backup: methods detail",
    ]
    assert slides[4].attrs.get("hidden") == "true"
    notes = [b for b in doc.blocks if isinstance(b, Paragraph) and b.attrs.get("slide_part") == "notes"]
    assert len(notes) == 3
    slide_ids = {s.id: s.index for s in slides}
    assert all(n.parent_id in slide_ids for n in notes)
    lst = next(b for b in doc.blocks if isinstance(b, ListBlock))
    assert lst.items[0].children[0].children[0].spans[0].text == "Great crested newt"
    tables = [b for b in doc.blocks if isinstance(b, Table)]
    merged = tables[0]
    assert merged.header_rows == 1 and any(c.col_span == 2 for c in merged.cells)
    chart = tables[1]
    assert chart.attrs["chart_type"] == "columnChart" and chart.caption
    assert [c.spans[0].text for c in chart.cells if c.row == 0] == ["Category", "Plants", "Animals"]
    assert all(b.provenance.source_page and b.provenance.path for b in doc.blocks)
    assert merged.provenance.bbox is not None and merged.provenance.bbox.page_width
    assert any(w.kind == WarningKind.HIDDEN_SLIDES_INCLUDED for w in doc.warnings)
    assert doc.metadata.title == "Pond Ecology 101" and doc.metadata.pages is None
    assert doc.metadata.slides == 5


def test_pptx_options(run: Run, fixture_bytes: Callable[[str], bytes]) -> None:
    doc = run(
        PptxConverter(), fixture_bytes("pptx-lecture"), "d.pptx", include_hidden_slides=False, include_notes=False
    )
    assert len([b for b in doc.blocks if isinstance(b, Slide)]) == 4
    assert not any(isinstance(b, Paragraph) and b.attrs.get("slide_part") for b in doc.blocks)
    capped = run(PptxConverter(), fixture_bytes("pptx-lecture"), "d.pptx", max_slides=2)
    assert len([b for b in capped.blocks if isinstance(b, Slide)]) == 2 and capped.truncated
    assert any(w.kind == WarningKind.SLIDE_CAP_REACHED for w in capped.warnings)


def test_pptx_picture_smartart_comments(run: Run) -> None:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
        "1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082"
    )
    pic = s.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(1))
    pic._element.nvPicPr.cNvPr.set("descr", "A heron")
    out = io.BytesIO()
    prs.save(out)
    doc = run(PptxConverter(), out.getvalue(), "p.pptx")
    img = next(b for b in doc.blocks if isinstance(b, Image))
    assert img.alt == "A heron" and img.ref.startswith("ppt/media/")


def test_xlsx_fixture(run: Run, fixture_bytes: Callable[[str], bytes]) -> None:
    doc = run(XlsxConverter(), fixture_bytes("xlsx-multi-sheet"), "w.xlsx")
    heads = [b for b in doc.blocks if isinstance(b, Heading)]
    assert [h.spans[0].text for h in heads] == ["Visits", "Summary", "Lookup", "Named ranges"]
    assert doc.metadata.sheets == ["Visits", "Summary", "Lookup"] and doc.metadata.pages is None
    assert heads[2].attrs["hidden"] == "true" and heads[2].attrs["sheet_state"] == "hidden"
    tables = [b for b in doc.blocks if isinstance(b, Table)]
    visits = tables[0]
    assert visits.column_types == ["text", "date", "int", "percent", "currency"]
    cells = {(c.row, c.col): c for c in visits.cells}
    assert cells[(4, 2)].formula == "=SUM(C2:C4)" and cells[(4, 2)].spans[0].text == "34"
    assert cells[(1, 1)].spans[0].text == "2025-04-03" and cells[(1, 3)].spans[0].text == "25.0%"
    assert visits.provenance.path == "Visits!A1:E5" and visits.caption
    assert tables[1].provenance.path == "Visits!A8:B9"
    summary = tables[2]
    assert summary.header_rows == 2 and any(c.col_span == 3 for c in summary.cells)
    comment = next(b for b in doc.blocks if isinstance(b, Comment))
    assert comment.anchor_text == "C2" and comment.author == "Carol Diaz" and comment.anchor_block_id == visits.id
    kinds = {w.kind for w in doc.warnings}
    assert {WarningKind.HIDDEN_SHEETS_INCLUDED, WarningKind.CELL_ERRORS} <= kinds
    assert WarningKind.FORMULAS_PRESENT not in kinds  # the renderer emits it once
    assert tables[-1].cells[2].spans[0].text == "SiteCodes"


def test_xlsx_options(run: Run, fixture_bytes: Callable[[str], bytes]) -> None:
    data = fixture_bytes("xlsx-multi-sheet")
    doc = run(XlsxConverter(), data, "w.xlsx", include_hidden_sheets=False, formulas="table")
    assert "Lookup" not in [b.spans[0].text for b in doc.blocks if isinstance(b, Heading)]
    assert any(isinstance(b, Table) and b.attrs.get("role") == "formulas" for b in doc.blocks)
    off = run(XlsxConverter(), data, "w.xlsx", formulas=False)
    assert not any(c.formula for b in off.blocks if isinstance(b, Table) for c in b.cells)
    capped = run(XlsxConverter(), data, "w.xlsx", max_rows=3)
    assert capped.truncated and any(w.kind == WarningKind.ROW_CAP_REACHED for w in capped.warnings)


def _workbook(build: Callable[[object], None]) -> bytes:
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    build(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_xlsx_uncalculated_and_serial_dates(run: Run) -> None:
    def build(wb: object) -> None:
        ws = wb.active  # type: ignore[attr-defined]
        ws.append(["Order date", "Amount"])
        ws.append([45000, 10])
        ws.append([45031, 20])
        ws["B4"] = "=SUM(B2:B3)"

    doc = run(XlsxConverter(), _workbook(build), "u.xlsx")
    kinds = {w.kind: w for w in doc.warnings}
    assert kinds[WarningKind.FORMULA_UNCALCULATED].count == 1
    assert "Sheet!A" in str(kinds[WarningKind.POSSIBLE_SERIAL_DATES].detail["columns"])
    table = next(b for b in doc.blocks if isinstance(b, Table))
    cell = next(c for c in table.cells if c.formula)
    assert cell.spans[0].text == "=SUM(B2:B3)"


def test_xlsx_empty_and_columns_cap(run: Run) -> None:
    doc = run(XlsxConverter(), _workbook(lambda wb: None), "e.xlsx")
    assert any(w.kind == WarningKind.EXTRACTION_EMPTY for w in doc.warnings)

    def wide(wb: object) -> None:
        wb.active.append(list(range(10)))  # type: ignore[attr-defined]

    capped = run(XlsxConverter(), _workbook(wide), "w.xlsx", max_cols=4)
    assert any(w.kind == WarningKind.COLUMNS_TRUNCATED for w in capped.warnings)


def test_cell_helpers() -> None:
    def c(v: object, fmt: str = "General", t: str = "n") -> XCell:
        return XCell(value=v, number_format=fmt, data_type=t, bold=False)

    assert fmt_value(c(True)) == ("TRUE", False)
    assert fmt_value(c(dt.datetime(2025, 1, 2, 3, 4))) == ("2025-01-02T03:04:00", False)
    assert fmt_value(c(dt.datetime(2025, 1, 2))) == ("2025-01-02", False)
    assert fmt_value(c(0.125, "0.00%")) == ("12.50%", False)
    assert fmt_value(c(3.0)) == ("3", False) and fmt_value(c(0.1)) == ("0.1", False)
    assert fmt_value(c("#DIV/0!", t="e")) == ("#DIV/0!", True)
    assert fmt_value(c(dt.timedelta(hours=1))) == ("1:00:00", False)
    grid: list[list[XCell | None]] = [[c("a"), c("b")], [c(1), c(2)], [], [], [c("x")]]
    assert regions(grid) == [(0, 0, 1, 1), (4, 0, 4, 0)]
    assert header_rows(grid[:2], False) == 1
    bold = [[XCell("h", "General", "s", True)], [XCell("v", "General", "s", False)]]
    assert header_rows(bold, True) == 1 and header_rows(bold, False) == 0
    assert col_letter(28) == "AB" and a1(0, 27) == "AB1"


@pytest.mark.parametrize(
    ("value", "fmt", "want"),
    [
        (120.5, '"$"#,##0.00', "$120.50"),
        (1234567.891, "#,##0.00", "1,234,567.89"),
        (-5, '"$"#,##0.00', "-$5.00"),
        (-5, "#,##0.00;(#,##0.00)", "(5.00)"),
        (0, "0.00;-0.00;[Red]zero", "zero"),
        (0.5, "0.0%", "50.0%"),
        (1500, "#,##0", "1,500"),
        (12345.678, "0.00E+00", "1.23E+04"),
        (9.5, "[$EUR-407] #,##0.00", "EUR 9.50"),
        (2.5, "0.0#", "2.5"),
        (42, "[Red]0", "42"),
        (1000, "#,##0_);(#,##0)", "1,000"),
        (3, "00", "03"),
        (2500000, "#,##0,", "2,500"),
        (7, "General", None),
        (0.25, "# ?/?", None),
        (45000, "yyyy-mm-dd", None),
    ],
)
def test_number_formats(value: float, fmt: str, want: str | None) -> None:
    from ezmd_converters.office.numfmt import format_number

    assert format_number(value, fmt) == want


def test_xlsx_header_hints(run: Run) -> None:
    def build(wb: object) -> None:
        from openpyxl.worksheet.table import Table as XTable

        ws = wb.active  # type: ignore[attr-defined]
        ws.title = "Codes"
        for row in (["Code", "Meaning"], ["NP", "North pond"], ["SP", "South pond"]):
            ws.append(row)
        ws.auto_filter.ref = "A1:B3"
        ws2 = wb.create_sheet("Staff")  # type: ignore[attr-defined]
        for row in (["Name", "Role"], ["Ann", "Lead"], ["Bo", "Field"]):
            ws2.append(row)
        ws2.add_table(XTable(displayName="StaffTable", ref="A1:B3"))
        ws3 = wb.create_sheet("Plain")  # type: ignore[attr-defined]
        for row in (["x", "y"], ["a", "b"]):
            ws3.append(row)

    doc = run(XlsxConverter(), _workbook(build), "h.xlsx")
    tables = [b for b in doc.blocks if isinstance(b, Table)]
    assert [t.header_rows for t in tables] == [1, 1, 0]


def test_xlsx_merged_region_keeps_empty_cells(run: Run, fixture_bytes: Callable[[str], bytes]) -> None:
    doc = run(XlsxConverter(), fixture_bytes("xlsx-merged-formulas"), "w.xlsx")
    shifts = next(b for b in doc.blocks if isinstance(b, Table))
    cells = {(c.row, c.col): c for c in shifts.cells}
    assert (cells[(0, 0)].col_span, cells[(2, 0)].row_span, cells[(4, 0)].row_span) == (4, 2, 2)
    # B7 is empty; inside a merged (HTML-rendered) table it must stay as a placeholder so 33 stays in column C.
    assert cells[(6, 1)].spans == [] and cells[(6, 2)].spans[0].text == "33"
    assert (3, 0) not in cells  # covered by the A3:A4 merge


def test_xlsx_sparse_region_without_merges_omits_empty_cells(run: Run, fixture_bytes: Callable[[str], bytes]) -> None:
    doc = run(XlsxConverter(), fixture_bytes("xlsx-uncalculated"), "w.xlsx")
    budget = next(b for b in doc.blocks if isinstance(b, Table))
    cells = {(c.row, c.col): c for c in budget.cells}
    assert (4, 2) not in cells and cells[(4, 3)].formula == "=SUM(D2:D4)"
    assert any(w.kind == WarningKind.FORMULA_UNCALCULATED for w in doc.warnings)
