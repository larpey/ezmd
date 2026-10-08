"""documents.xlsx: spreadsheet converter on openpyxl (MIT) over the sanitized package (Part 2 2c 24-34).

Every worksheet in workbook order, hidden and veryHidden included (sheet Heading attrs hidden="true" plus
`hidden_sheets_included`), one Table per contiguous region, cached values in the cell spans, formula text in
`TableCell.formula`, `Table.column_types`, merged cells as spans, cell comments, named ranges. A second
`data_only=False` pass reads formulas. Small workbooks load fully (merges, comments, bold headers); large ones
stream in read-only mode.
"""

from __future__ import annotations

import io
import logging
import warnings as pywarnings
from typing import Any

from intomd.context import Limits
from intomd.core.textclean import CleanStats, clean_text
from intomd.inputs import InputRef
from intomd.ir import (
    Comment,
    Document,
    Heading,
    InlineSpan,
    Metadata,
    Provenance,
    SourceType,
    Table,
    TableCell,
    Warning,
    WarningKind,
)
from intomd.registry import ConversionError, ConvertOptions
from intomd_converters.office._common import (
    FAMILY,
    MB,
    XLSX_EXTS,
    XLSX_MIMES,
    BlockList,
    OfficeOptions,
    check_size,
    clean_warning,
    confidence,
)
from intomd_converters.office._package import OfficePackage
from intomd_converters.office.xlsx_cells import (
    DATE_HEADER,
    XCell,
    a1,
    column_types,
    fmt_value,
    header_rows,
    regions,
)

log = logging.getLogger(__name__)

FULL_LOAD_BYTES = 20 * MB
"""Uncompressed worksheet XML up to this size loads in normal mode (merges, comments, fonts)."""
EMPTY_ROW_STOP = 50


class SheetStats:
    def __init__(self) -> None:
        self.hidden: list[str] = []
        self.row_capped: list[str] = []
        self.col_capped: list[str] = []
        self.errors = 0
        self.uncalculated = 0
        self.formulas = 0
        self.serial_dates: list[str] = []


class XlsxConverter:
    id = "documents.xlsx"
    family = FAMILY
    priority = 10
    experimental = False
    requires_extras: tuple[str, ...] = ()
    mimes: tuple[str, ...] = XLSX_MIMES
    limits = Limits(max_bytes=200 * MB, max_rows=10_000, max_cols=256, max_sheets=50, timeout_s=120)

    def can_handle(self, ref: InputRef) -> float:
        return confidence(ref, XLSX_MIMES, XLSX_EXTS)

    def convert(self, ref: InputRef, options: ConvertOptions) -> Document:
        check_size(ref, self.limits.max_bytes or 200 * MB, "Excel")
        pkg = OfficePackage(ref.read(), what="XLSX")
        return convert_package(pkg, ref.display, options, mime=ref.detected.mime if ref.detected else None)


def _load(data: bytes, *, read_only: bool, data_only: bool) -> Any:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    with pywarnings.catch_warnings():
        pywarnings.simplefilter("ignore")
        try:
            return load_workbook(io.BytesIO(data), read_only=read_only, data_only=data_only, keep_links=False)
        except Exception as e:  # openpyxl raises many types on malformed packages
            raise ConversionError(
                f"openpyxl could not open the workbook: {e}", user_message="The workbook is damaged."
            ) from e


def convert_package(pkg: OfficePackage, source: str, options: ConvertOptions, *, mime: str | None = None) -> Document:
    opts = OfficeOptions.from_options(options)
    sheet_bytes = sum(pkg.members[n].file_size for n in pkg.names() if n.startswith("xl/worksheets/"))
    full = sheet_bytes <= FULL_LOAD_BYTES
    sane = pkg.sanitized_bytes()
    wb = _load(sane, read_only=not full, data_only=True)
    wbf = _load(sane, read_only=not full, data_only=False) if opts.formulas != "off" else None
    try:
        doc = _convert(wb, wbf, full, source, opts, options, mime)
    finally:
        for w in (wb, wbf):
            if w is not None and hasattr(w, "close"):
                w.close()
    doc.warnings[:0] = pkg.warnings()
    return doc.finalize()


def _convert(
    wb: Any, wbf: Any, full: bool, source: str, opts: OfficeOptions, options: ConvertOptions, mime: str | None
) -> Document:
    props = wb.properties
    meta = Metadata(
        source=source,
        source_type=SourceType.XLSX,
        mime=mime,
        title=(props.title or None) if props is not None else None,
        author=(props.creator or None) if props is not None else None,
        authors=[props.creator] if props is not None and props.creator else [],
    )
    doc = Document(metadata=meta)
    out = BlockList()
    stats = CleanStats()
    st = SheetStats()
    sheets = list(wb.worksheets)
    meta.sheets = [str(ws.title) for ws in sheets]
    for idx, ws in enumerate(sheets[: opts.max_sheets], start=1):
        options.ctx.check_deadline()
        options.ctx.progress("sheets", idx / max(len(sheets), 1), ws.title)
        state = getattr(ws, "sheet_state", "visible") or "visible"
        if state != "visible":
            if not opts.include_hidden_sheets:
                continue
            st.hidden.append(ws.title)
        fws = wbf[ws.title] if wbf is not None and ws.title in wbf.sheetnames else None
        hints = header_hints(wb, ws)
        _sheet(ws, fws, idx, full, source, opts, options, out, stats, st, state, hints)
        doc.blocks = list(out.blocks)
        options.ctx.publish_partial(doc)
    _named_ranges(wb, out, source, stats)
    doc.blocks = list(out.blocks)
    doc.warnings.extend(sheet_warnings(st, len(sheets), opts))
    cw = clean_warning(stats)
    if cw is not None:
        doc.warnings.append(cw)
    if st.row_capped or st.col_capped or len(sheets) > opts.max_sheets:
        doc.truncated = True
    if not any(isinstance(b, Table) for b in doc.blocks):
        doc.warnings.append(Warning(kind=WarningKind.EXTRACTION_EMPTY, message="The workbook has no cell values."))
    return doc


def _read_grid(ws: Any, fws: Any, opts: OfficeOptions, full: bool) -> tuple[list[list[XCell | None]], bool, bool]:
    grid: list[list[XCell | None]] = []
    empty_run = 0
    row_capped = col_capped = False
    formulas: dict[tuple[int, int], str] = {}
    if fws is not None:
        for r, row in enumerate(fws.iter_rows(min_row=1, min_col=1, max_col=opts.max_cols)):
            if r >= opts.max_rows:
                break
            for c, cell in enumerate(row):
                f = _formula_text(getattr(cell, "value", None))
                if f:
                    formulas[(r, c)] = f
    for r, row in enumerate(ws.iter_rows(min_row=1, min_col=1, max_col=opts.max_cols + 1)):
        if r >= opts.max_rows:
            row_capped = any(getattr(c, "value", None) is not None for c in row) or row_capped
            if row_capped:
                break
            continue
        cells: list[XCell | None] = []
        for c, cell in enumerate(row):
            if c >= opts.max_cols:
                if getattr(cell, "value", None) is not None:
                    col_capped = True
                break
            value = getattr(cell, "value", None)
            f = formulas.get((r, c))
            if value is None and f is None:
                cells.append(None)
                continue
            bold = bool(full and getattr(getattr(cell, "font", None), "b", False))
            cells.append(
                XCell(
                    value=value,
                    number_format=str(getattr(cell, "number_format", "General") or "General"),
                    data_type=str(getattr(cell, "data_type", "n") or "n"),
                    bold=bold,
                    formula=f,
                )
            )
        while cells and cells[-1] is None:
            cells.pop()
        grid.append(cells)
        empty_run = 0 if cells else empty_run + 1
        if empty_run >= EMPTY_ROW_STOP:
            break
    while grid and not grid[-1]:
        grid.pop()
    return grid, row_capped, col_capped


def _formula_text(value: object) -> str | None:
    if isinstance(value, str) and value.startswith("=") and len(value) > 1:
        return value
    text = getattr(value, "text", None)  # openpyxl ArrayFormula / DataTableFormula
    if isinstance(text, str) and text:
        return text if text.startswith("=") else "=" + text
    return None


def _sheet(
    ws: Any,
    fws: Any,
    idx: int,
    full: bool,
    source: str,
    opts: OfficeOptions,
    options: ConvertOptions,
    out: BlockList,
    stats: CleanStats,
    st: SheetStats,
    state: str,
    hints: set[tuple[int, int, int]] | None = None,
) -> None:
    title = str(ws.title)
    attrs = {"sheet_state": state}
    if state != "visible":
        attrs["hidden"] = "true"
    heading_id = out.add(
        Heading(
            level=2,
            spans=[InlineSpan(text=clean_text(title, stats))],
            provenance=Provenance(source=source, source_page=idx, source_label=title, path=f"{title}!A1"),
            attrs=attrs,
        )
    )
    grid, row_capped, col_capped = _read_grid(ws, fws, opts, full)
    if row_capped:
        st.row_capped.append(title)
    if col_capped:
        st.col_capped.append(title)
    merges = _merges(ws) if full else []
    regs = regions(grid)
    table_ids: list[tuple[tuple[int, int, int, int], str]] = []
    for n, reg in enumerate(regs, start=1):
        options.ctx.check_deadline()
        table = sheet_table(grid, reg, merges, title, idx, source, stats, st, full, len(regs) > 1, n, hints=hints)
        if table is not None:
            tid = out.add(table)
            table_ids.append((reg, tid))
            if opts.formulas == "table" and any(c.formula for c in table.cells):
                out.add(_formula_table(table))
    if full and opts.comments:
        for c in _comments(ws, title, idx, source, stats, table_ids, heading_id):
            out.add(c)


def _merges(ws: Any) -> list[tuple[int, int, int, int]]:
    out: list[tuple[int, int, int, int]] = []
    for rng in getattr(getattr(ws, "merged_cells", None), "ranges", []) or []:
        out.append((rng.min_row - 1, rng.min_col - 1, rng.max_row - 1, rng.max_col - 1))
    return out


def sheet_table(
    grid: list[list[XCell | None]],
    reg: tuple[int, int, int, int],
    merges: list[tuple[int, int, int, int]],
    sheet: str,
    idx: int,
    source: str,
    stats: CleanStats,
    st: SheetStats,
    bold_known: bool,
    many: bool,
    n: int,
    *,
    hints: set[tuple[int, int, int]] | None = None,
) -> Table | None:
    r0, c0, r1, c1 = reg
    n_rows, n_cols = r1 - r0 + 1, c1 - c0 + 1
    rows = [[(grid[r][c] if c < len(grid[r]) else None) for c in range(c0, c1 + 1)] for r in range(r0, r1 + 1)]
    covered: set[tuple[int, int]] = set()
    spans: dict[tuple[int, int], tuple[int, int]] = {}
    for mr0, mc0, mr1, mc1 in merges:
        if not (r0 <= mr0 <= r1 and c0 <= mc0 <= c1):
            continue
        er, ec = min(mr1, r1), min(mc1, c1)
        spans[(mr0 - r0, mc0 - c0)] = (er - mr0 + 1, ec - mc0 + 1)
        for rr in range(mr0, er + 1):
            for cc in range(mc0, ec + 1):
                if (rr, cc) != (mr0, mc0):
                    covered.add((rr - r0, cc - c0))
    hdr = header_rows(rows, bold_known)
    if hdr == 0 and n_rows > 1 and _hinted(hints, r0, c0, c1):
        hdr = 1
    cells: list[TableCell] = []
    for ri, row in enumerate(rows):
        for ci, xc in enumerate(row):
            if (ri, ci) in covered:
                continue
            rs, cs = spans.get((ri, ci), (1, 1))
            text = ""
            formula = None
            raw = None
            if xc is not None:
                text, is_error = fmt_value(xc)
                st.errors += int(is_error)
                formula = xc.formula
                if formula:
                    st.formulas += 1
                    if xc.value is None:
                        st.uncalculated += 1
                        text = formula
                if isinstance(xc.value, int | float) and not isinstance(xc.value, bool):
                    raw = repr(xc.value)
            if not text and (rs, cs) == (1, 1) and formula is None and not spans:
                # Sparse cells are fine without merges; with merges the table renders as HTML, where a
                # missing cell would shift every later cell in its row one column left.
                continue
            cells.append(
                TableCell(
                    spans=[InlineSpan(text=clean_text(text, stats))] if text else [],
                    row=ri,
                    col=ci,
                    row_span=rs,
                    col_span=cs,
                    is_header=ri < hdr,
                    formula=formula,
                    raw_value=raw,
                )
            )
    if not cells:
        return None
    types = column_types(rows[hdr:], n_cols)
    _serial_dates(rows, hdr, sheet, c0, st)
    ref = f"{a1(r0, c0)}:{a1(r1, c1)}"
    attrs = {"sheet": sheet, "region": ref}
    if n_cols > 6:
        attrs["wide"] = "true"
    return Table(
        cells=cells,
        n_rows=n_rows,
        n_cols=n_cols,
        header_rows=hdr,
        column_types=types,
        caption=[InlineSpan(text=f"Region {n} ({ref})")] if many else None,
        provenance=Provenance(source=source, source_page=idx, source_label=sheet, path=f"{sheet}!{ref}"),
        attrs=attrs,
    )


def _serial_dates(rows: list[list[XCell | None]], hdr: int, sheet: str, c0: int, st: SheetStats) -> None:
    if hdr == 0:
        return
    for ci, head in enumerate(rows[hdr - 1]):
        if head is None or not isinstance(head.value, str) or not DATE_HEADER.search(head.value):
            continue
        vals = [r[ci] for r in rows[hdr:] if ci < len(r) and r[ci] is not None]
        nums = [v for v in vals if isinstance(v.value, int | float) and not isinstance(v.value, bool)]  # type: ignore[union-attr]
        if (
            nums
            and len(nums) == len(vals)
            and all(
                v.number_format == "General" and 20_000 <= float(v.value) <= 80_000 and float(v.value).is_integer()  # type: ignore[union-attr, arg-type]
                for v in nums
            )
        ):
            st.serial_dates.append(f"{sheet}!{a1(0, c0 + ci)[:-1]}")


def header_hints(wb: Any, ws: Any) -> set[tuple[int, int, int]]:
    """Rows the workbook itself marks as headers, as (row, col0, col1), 0-based: the row just above a defined
    name's range on this sheet (the name covers the data, not its labels), and the first row of an autofilter
    or a table with a header row."""
    from openpyxl.utils.cell import range_boundaries  # type: ignore[import-untyped]

    out: set[tuple[int, int, int]] = set()

    def add(ref: str, offset: int) -> None:
        try:
            min_col, min_row, max_col, _max_row = range_boundaries(ref.replace("$", ""))
        except (ValueError, TypeError):
            return
        if min_col is None or min_row is None or max_col is None:
            return
        row = min_row - 1 - offset
        if row >= 0:
            out.add((row, min_col - 1, max_col - 1))

    names = getattr(wb, "defined_names", None)
    for _name, dn in list(names.items()) if names is not None and hasattr(names, "items") else []:
        try:
            dests = list(dn.destinations)
        except Exception as e:  # malformed or external references
            log.debug("skipping defined name %s: %s", _name, e)
            continue
        for sheet, ref in dests:
            if sheet == ws.title:
                add(str(ref), 1)
    filt = getattr(getattr(ws, "auto_filter", None), "ref", None)
    if filt:
        add(str(filt), 0)
    tables = getattr(ws, "tables", None)
    for tbl in tables.values() if tables is not None and hasattr(tables, "values") else []:
        if getattr(tbl, "headerRowCount", 1) and getattr(tbl, "ref", None):
            add(str(tbl.ref), 0)
    return out


def _hinted(hints: set[tuple[int, int, int]] | None, r0: int, c0: int, c1: int) -> bool:
    return bool(hints) and any(row == r0 and a <= c1 and c0 <= b for row, a, b in hints or ())


def _formula_table(table: Table) -> Table:
    cells = [
        TableCell(
            spans=[InlineSpan(text=c.formula or "")] if c.formula else [],
            row=c.row,
            col=c.col,
            row_span=c.row_span,
            col_span=c.col_span,
            is_header=c.is_header,
        )
        for c in table.cells
    ]
    return Table(
        cells=cells,
        n_rows=table.n_rows,
        n_cols=table.n_cols,
        header_rows=table.header_rows,
        caption=[InlineSpan(text="Formulas")],
        provenance=table.provenance,
        attrs={**table.attrs, "role": "formulas"},
    )


def _comments(
    ws: Any,
    sheet: str,
    idx: int,
    source: str,
    stats: CleanStats,
    tables: list[tuple[tuple[int, int, int, int], str]],
    heading_id: str,
) -> list[Comment]:
    out: list[Comment] = []
    for row in ws.iter_rows():
        for cell in row:
            cm = getattr(cell, "comment", None)
            if cm is None or not getattr(cm, "text", ""):
                continue
            r, c = cell.row - 1, cell.column - 1
            anchor = next((tid for (r0, c0, r1, c1), tid in tables if r0 <= r <= r1 and c0 <= c <= c1), heading_id)
            out.append(
                Comment(
                    author=(cm.author or None),
                    spans=[InlineSpan(text=clean_text(str(cm.text).strip(), stats))],
                    anchor_block_id=anchor,
                    anchor_text=cell.coordinate,
                    provenance=Provenance(
                        source=source, source_page=idx, source_label=sheet, path=f"{sheet}!{cell.coordinate}"
                    ),
                )
            )
    return out


def _named_ranges(wb: Any, out: BlockList, source: str, stats: CleanStats) -> None:
    names: list[tuple[str, str]] = []
    dn = getattr(wb, "defined_names", None)
    items = list(dn.items()) if dn is not None and hasattr(dn, "items") else []
    for name, d in items:
        names.append((str(name), str(getattr(d, "attr_text", "") or "")))
    for ws in wb.worksheets:
        sdn = getattr(ws, "defined_names", None)
        if sdn is not None and hasattr(sdn, "items"):
            for name, d in sdn.items():
                names.append((f"{ws.title}!{name}", str(getattr(d, "attr_text", "") or "")))
    names = [(n, v) for n, v in names if not n.startswith("_xlnm.")]
    if not names:
        return
    prov = Provenance(source=source, path="definedNames")
    out.add(Heading(level=2, spans=[InlineSpan(text="Named ranges")], provenance=prov, attrs={"role": "named_ranges"}))
    cells = [
        TableCell(spans=[InlineSpan(text="Name")], row=0, col=0, is_header=True),
        TableCell(spans=[InlineSpan(text="Refers to")], row=0, col=1, is_header=True),
    ]
    for i, (n, v) in enumerate(sorted(names), start=1):
        cells.append(TableCell(spans=[InlineSpan(text=clean_text(n, stats))], row=i, col=0))
        cells.append(TableCell(spans=[InlineSpan(text=clean_text(v, stats))] if v else [], row=i, col=1))
    out.add(
        Table(
            cells=cells, n_rows=len(names) + 1, n_cols=2, header_rows=1, provenance=prov, column_types=["text", "text"]
        )
    )


def sheet_warnings(st: SheetStats, n_sheets: int, opts: OfficeOptions) -> list[Warning]:
    out: list[Warning] = []
    if st.hidden:
        out.append(
            Warning(
                kind=WarningKind.HIDDEN_SHEETS_INCLUDED,
                message=f"Included hidden sheets: {', '.join(st.hidden)}.",
                count=len(st.hidden),
                detail={"sheets": ",".join(st.hidden)},
            )
        )
    # formulas_present is emitted by the renderer whenever a table cell carries a formula (once per output).
    if st.uncalculated:
        out.append(
            Warning(
                kind=WarningKind.FORMULA_UNCALCULATED,
                message=f"{st.uncalculated} formula cells have no cached value; the formula text is shown.",
                count=st.uncalculated,
            )
        )
    if st.errors:
        out.append(
            Warning(kind=WarningKind.CELL_ERRORS, message=f"{st.errors} cells contain error values.", count=st.errors)
        )
    if st.serial_dates:
        out.append(
            Warning(
                kind=WarningKind.POSSIBLE_SERIAL_DATES,
                message=f"Columns that look like unformatted date serials: {', '.join(st.serial_dates)}.",
                count=len(st.serial_dates),
                detail={"columns": ",".join(st.serial_dates)},
            )
        )
    for sheet in st.row_capped:
        out.append(
            Warning(
                kind=WarningKind.ROW_CAP_REACHED,
                message=f"Sheet {sheet!r} was cut at {opts.max_rows} rows; raise office.max_rows to convert more.",
                detail={"sheet": sheet, "max_rows": opts.max_rows},
            )
        )
    for sheet in st.col_capped:
        out.append(
            Warning(
                kind=WarningKind.COLUMNS_TRUNCATED,
                message=f"Sheet {sheet!r} was cut at {opts.max_cols} columns; raise office.max_cols to convert more.",
                detail={"sheet": sheet, "max_cols": opts.max_cols},
            )
        )
    if n_sheets > opts.max_sheets:
        out.append(
            Warning(
                kind=WarningKind.TRUNCATED,
                message=f"Converted the first {opts.max_sheets} of {n_sheets} sheets; raise office.max_sheets.",
                detail={"reason": "sheet_cap", "sheets": n_sheets, "cap": opts.max_sheets},
            )
        )
    return out
